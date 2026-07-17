"""Document loader.

Detects the evidence type and opens the resources needed downstream:

* PDF  -> a PyMuPDF document used for both rendering and structure analysis.
* DOCX -> a PyMuPDF document (opened directly) for rendering, a converted
          in-memory PDF for structure analysis, and a python-docx handle for
          formatting/font extraction.
* Raster image -> a Pillow-backed document which preserves the submitted
          pixels for visual analysis. Multi-frame TIFF files expose one page
          per frame. PDF-only structure and font analysis do not apply.

All loading is local and deterministic.
"""

from __future__ import annotations

import os
from dataclasses import dataclass, field
from typing import Optional

import fitz  # PyMuPDF
from PIL import Image

from ..utils.config import Config
from ..utils.logging import get_logger

logger = get_logger("loader")

_PDF_EXTS = {".pdf"}
_DOCX_EXTS = {".docx"}
_RASTER_EXT_TO_TYPE = {
    ".png": "png",
    ".jpg": "jpeg",
    ".jpeg": "jpeg",
    ".tif": "tiff",
    ".tiff": "tiff",
    ".bmp": "bmp",
    ".webp": "webp",
    ".gif": "gif",
}


class RasterDocument:
    """Small Pillow-backed page source used by :func:`render_document`.

    Pillow keeps image decoding lazy, which avoids loading every frame of a
    potentially large TIFF into memory at once. Callers obtain independent
    frame copies through :meth:`get_frame` so seeking the source is safe.
    """

    is_raster_document = True

    def __init__(self, path: str) -> None:
        self.path = path
        self._image = Image.open(path)
        self.page_count = int(getattr(self._image, "n_frames", 1) or 1)

    def get_frame(self, index: int) -> tuple["Image.Image", dict]:
        if index < 0 or index >= self.page_count:
            raise IndexError(index)
        self._image.seek(index)
        # Capture frame-specific metadata before copying; Pillow may update
        # ``info`` while seeking through a multi-frame TIFF.
        info = dict(self._image.info)
        return self._image.copy(), info

    def close(self) -> None:
        self._image.close()


@dataclass
class LoadedDocument:
    """Container for everything the pipeline needs from one input file."""

    path: str
    doc_type: str  # "pdf", "docx", or a normalized raster format name
    render_doc: object  # fitz.Document or RasterDocument
    pdf_doc: Optional["fitz.Document"] = None  # PDF view for structure analysis
    docx_obj: object = None  # python-docx Document, for DOCX only
    warnings: list[str] = field(default_factory=list)

    def close(self) -> None:
        for doc in (self.render_doc, self.pdf_doc):
            try:
                if doc is not None:
                    doc.close()
            except Exception:  # pragma: no cover - defensive
                pass


def detect_type(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()

    # Prefer strong content signatures. This both supports extensionless
    # evidence and reports the actual encoding for mislabeled raster files.
    try:
        with open(path, "rb") as fh:
            head = fh.read(16)
        if head.startswith(b"%PDF"):
            return "pdf"
        if head.startswith(b"\x89PNG\r\n\x1a\n"):
            return "png"
        if head.startswith(b"\xff\xd8\xff"):
            return "jpeg"
        if head.startswith((b"II*\x00", b"MM\x00*")):
            return "tiff"
        if head.startswith(b"BM"):
            return "bmp"
        if head.startswith(b"RIFF") and head[8:12] == b"WEBP":
            return "webp"
        if head.startswith((b"GIF87a", b"GIF89a")):
            return "gif"
        if head.startswith(b"PK\x03\x04") and ext in _DOCX_EXTS:
            return "docx"
    except OSError:
        pass

    # Extensions remain useful when a format lacks an unambiguous short
    # signature, while the actual decoder still validates the file on load.
    if ext in _PDF_EXTS:
        return "pdf"
    if ext in _DOCX_EXTS:
        return "docx"
    if ext in _RASTER_EXT_TO_TYPE:
        return _RASTER_EXT_TO_TYPE[ext]
    raise ValueError(f"Unsupported or unrecognised document type: {path!r}")


def load_document(path: str, config: Config) -> LoadedDocument:
    """Open ``path`` and return a :class:`LoadedDocument`.

    Raises ``FileNotFoundError`` / ``ValueError`` for invalid inputs.
    """
    if not os.path.exists(path):
        raise FileNotFoundError(f"Document not found: {path!r}")
    if not os.path.isfile(path):
        raise ValueError(f"Not a file: {path!r}")

    doc_type = detect_type(path)
    logger.info("Loading %s document: %s", doc_type, path)

    if doc_type == "pdf":
        render_doc = fitz.open(path)
        return LoadedDocument(
            path=path,
            doc_type="pdf",
            render_doc=render_doc,
            pdf_doc=render_doc,  # same handle; structure analyzer reads only
        )

    if doc_type in set(_RASTER_EXT_TO_TYPE.values()):
        try:
            render_doc = RasterDocument(path)
        except Exception as exc:
            raise ValueError(f"Unable to decode {doc_type.upper()} image: {path!r}") from exc
        return LoadedDocument(
            path=path,
            doc_type=doc_type,
            render_doc=render_doc,
            # Raw images have no PDF object/font structure to inspect.
            pdf_doc=None,
        )

    # DOCX
    render_doc = fitz.open(path)
    loaded = LoadedDocument(path=path, doc_type="docx", render_doc=render_doc)

    # Converted PDF view enables PDF-structure analysis (images, fonts, objects).
    if config.get("features.enable_pdf_analysis", True):
        try:
            pdf_bytes = render_doc.convert_to_pdf()
            loaded.pdf_doc = fitz.open("pdf", pdf_bytes)
        except Exception as exc:  # pragma: no cover - depends on libmupdf build
            loaded.warnings.append(
                f"DOCX->PDF conversion failed; structure analysis limited: {exc}"
            )
            logger.warning("DOCX->PDF conversion failed: %s", exc)

    # python-docx handle for native formatting/font extraction.
    try:
        import docx  # local import keeps PDF path lightweight

        loaded.docx_obj = docx.Document(path)
    except Exception as exc:  # pragma: no cover - defensive
        loaded.warnings.append(f"python-docx load failed: {exc}")
        logger.warning("python-docx load failed: %s", exc)

    return loaded
